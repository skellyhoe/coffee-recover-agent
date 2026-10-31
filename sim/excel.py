"""The vendor's Excel export: the independent source, with product names but no codes."""
from openpyxl import Workbook, load_workbook

HEADER = ["order_id", "date", "time", "product_name", "price"]


def write_export(world, path):
    """One row per cup. A cutoff (day, time) models an export taken midway through that day."""
    cutoff = world.get("export_cutoff")
    wb = Workbook()
    ws = wb.active
    ws.title = "cups"
    ws.append(HEADER)
    for o in world["orders"]:
        if cutoff and o["date"] == cutoff["date"] and o["time"] >= cutoff["time"]:
            continue
        for c in o["cups"]:
            ws.append([o["id"], o["date"], o["time"], c["true_name"], c["price"]])
    meta = wb.create_sheet("meta")
    last = cutoff or {"date": world["days"][-1], "time": "23:59"}
    meta.append(["exported_date", last["date"]])
    meta.append(["exported_time", last["time"]])
    wb.save(path)


def read_export(path):
    wb = load_workbook(path, read_only=True)
    rows = list(wb["cups"].iter_rows(min_row=2, values_only=True))
    meta = dict(wb["meta"].iter_rows(values_only=True))
    wb.close()
    return {
        "rows": [dict(zip(HEADER, r)) for r in rows],
        "exported_date": meta["exported_date"],
        "exported_time": meta["exported_time"],
    }
